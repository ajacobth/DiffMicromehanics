"""service_pipeline.py — batch full-pipeline runner (Stages 1 → 2 → 3 from Excel).

Reads a measurements-only .xlsx file, runs elastic → thermoelastic → thermal inverse
in sequence entirely in memory, and returns a result dict containing both the summary
and the full save_data needed to persist to the DB later.

Nothing is saved to the database here — saving is delegated to the caller
(agent_tools.save_to_card) after the user has confirmed a card name.

Stage 2 and 3 inputs are built from Stage 1 in-memory results — no DB round-trips
between stages are required.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

import numpy as np

# ── Constants (mirror agent_tools.py exactly) ─────────────────────────────────

_STAGE1_FREE = ["a11", "a22", "fiber_massfrac", "ar", "matrix_modulus", "matrix_poisson"]
_STAGE1_FREE_WITHOUT_POISSON = ["a11", "a22", "fiber_massfrac", "ar", "matrix_modulus"]
_SHEAR_POISSON_MEASUREMENTS = frozenset(["G12", "G13", "G23", "nu12", "nu13", "nu23"])
_TE_FREE = ["f_cte1", "f_cte2", "m_cte"]
_TE_INIT = {"f_cte1": 1.5e-6, "f_cte2": 17.5e-6, "m_cte": 75.0e-6}

_DEFAULT_BOUNDS: dict = {
    "a11":            (0.50,   0.85),
    "a22":            (0.01,   0.40),
    "a12":            (-0.10,  0.10),
    "a13":            (-0.10,  0.10),
    "a23":            (-0.10,  0.10),
    "fiber_massfrac": (0.05,   0.60),
    "ar":             (5.0,  100.0),
    "matrix_modulus": (2000.0, 5000.0),
    "matrix_poisson": (0.33,   0.42),
    "f_cte1":         (-2e-6,  5e-6),
    "f_cte2":         (5e-6,  30e-6),
    "m_cte":          (30e-6, 120e-6),
}

_ELASTIC_SOLVER_CFG: dict = {
    "method":             "lbfgsb",
    "constraint_penalty": 10000.0,
    "use_epsilon_loss":   False,
    "epsilon_scale":      0.5,
    "maxiter":            300,
    "tol":                1e-6,
    "seed":               42,
}

_PROBLEM_JSON_INIT: dict = {
    "a11": 0.6, "a22": 0.1, "a12": 0.0, "a13": 0.0, "a23": 0.0,
    "fiber_massfrac": 0.20, "ar": 20.0,
}

_ELASTIC_TO_MODEL: dict = {
    "E1_MPa": "E1", "E2_MPa": "E2", "E3_MPa": "E3",
    "G12_MPa": "G12", "G13_MPa": "G13", "G23_MPa": "G23",
    "nu12": "nu12", "nu13": "nu13", "nu23": "nu23",
}
_ELASTIC_SIGMA_TO_MODEL: dict = {
    "E1_sigma_MPa": "E1", "E2_sigma_MPa": "E2", "E3_sigma_MPa": "E3",
    "G12_sigma_MPa": "G12", "G13_sigma_MPa": "G13", "G23_sigma_MPa": "G23",
    "nu12_sigma": "nu12", "nu13_sigma": "nu13", "nu23_sigma": "nu23",
}
_CTE_TO_MODEL: dict = {
    "CTE11_per_K": "CTE11",
    "CTE22_per_K": "CTE22",
    "CTE33_per_K": "CTE33",
}
_CTE_SIGMA_TO_MODEL: dict = {
    "CTE11_sigma_per_K": "CTE11",
    "CTE22_sigma_per_K": "CTE22",
    "CTE33_sigma_per_K": "CTE33",
}

_STAGE1_POOR_FIT_THRESHOLD = 0.05
_STAGE2_POOR_FIT_THRESHOLD = 0.05
_STAGE3_POOR_FIT_THRESHOLD = 1e-2


# ── Excel readers ─────────────────────────────────────────────────────────────

def _read_measurements_sheet(wb) -> dict[str, float]:
    """Read key-value pairs from the 'measurements' sheet (col A = key, col B = value).

    Rows where column A starts with '#' or is blank are skipped.
    """
    sheet_names = {s.lower(): s for s in wb.sheetnames}
    sheet_name = sheet_names.get("measurements")
    if sheet_name is None:
        raise ValueError(
            f"Sheet 'measurements' not found. Available: {list(wb.sheetnames)}"
        )
    ws = wb[sheet_name]
    data: dict[str, float] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        if not key or key.startswith("#"):
            continue
        if len(row) < 2 or row[1] is None:
            continue
        try:
            data[key] = float(row[1])
        except (TypeError, ValueError):
            pass
    return data


def _read_thermal_sheet(wb) -> tuple[Optional[np.ndarray], Optional[dict]]:
    """Read k vs T data from the 'thermal' sheet.

    Returns (temperatures_°C, K_data) or (None, None) if absent/empty.
    K_data = {"K11": array|None, "K22": array|None, "K33": array|None}
    """
    sheet_names = {s.lower(): s for s in wb.sheetnames}
    sheet_name = sheet_names.get("thermal")
    if sheet_name is None:
        return None, None

    ws = wb[sheet_name]
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return None, None

    # Skip comment/blank rows to find the header
    rows = []
    header_found = False
    for row in all_rows:
        first = str(row[0]).strip() if row and row[0] is not None else ""
        if not header_found:
            if first.startswith("#") or not first:
                continue
            header_found = True
        rows.append(row)

    if len(rows) < 2:
        return None, None

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]

    _T_ALIASES = {"temperature_c", "temperature", "t", "temp_c", "temp"}
    _K_ALIASES = {
        "K11": {"k11_wmk", "k11"},
        "K22": {"k22_wmk", "k22"},
        "K33": {"k33_wmk", "k33"},
    }

    t_idx = next((i for i, h in enumerate(header) if h in _T_ALIASES), None)
    if t_idx is None:
        raise ValueError(
            f"No temperature column in 'thermal' sheet. "
            f"Expected one of: temperature_c, temperature, T. Found: {header}"
        )

    k_idx: dict[str, Optional[int]] = {}
    for key, aliases in _K_ALIASES.items():
        k_idx[key] = next((i for i, h in enumerate(header) if h in aliases), None)

    if k_idx["K11"] is None:
        raise ValueError("K11 column required in 'thermal' sheet but not found.")

    data_rows = [r for r in rows[1:] if r and r[t_idx] is not None]
    if not data_rows:
        return None, None

    temps = np.array([float(r[t_idx]) for r in data_rows])
    K_data: dict[str, Optional[np.ndarray]] = {}
    for key, idx in k_idx.items():
        if idx is not None:
            try:
                K_data[key] = np.array([float(r[idx]) for r in data_rows])
            except (TypeError, ValueError):
                K_data[key] = None
        else:
            K_data[key] = None

    return temps, K_data


# ── Main pipeline ─────────────────────────────────────────────────────────────

def run_pipeline(
    file_path: str,
    fiber_id: int,
    polymer_id: int,
    printer_id: int,
    fiber_massfrac: float = -1.0,
    aspect_ratio: float = -1.0,
    a11: float = -1.0,
    a22: float = -1.0,
    n_restarts: int = 20,
) -> dict:
    """Run Stages 1 → 2 → 3 from an Excel measurements file. Does NOT save to DB.

    Returns a result dict with per-stage status + save_data for later persistence:
    {
        "stage1": {
            "status":    "pass"|"fail"|"skipped",
            "fit_error": float | None,
            "inferred":  dict,       # opt_free values
            "save_data": dict | None # passed to save_inverse_result when saving
        },
        "stage2": { ... same ... },
        "stage3": {
            "status":    ...,
            "fit_error": float | None,
            "inferred":  dict,
            "save_data": dict | None  # passed to save_thermal_result when saving
        },
        "errors": [str],
    }

    Abort rules:
      Stage 1 fail  → abort; stage2/stage3 status = "skipped"
      No CTE data   → stage2/stage3 status = "skipped"
      Stage 2 fail  → abort stage3; stage3 status = "skipped"
      No thermal sheet → stage3 status = "skipped"
    """
    import core.services.service_material as _smat
    import core.services.service_forward as _sfwd
    import core.services.service_inverse as _sinv
    import core.inverse_thermal as _ithermal

    def _skipped(note: str = "") -> dict:
        d: dict = {"status": "skipped", "fit_error": None, "inferred": {}, "save_data": None}
        if note:
            d["note"] = note
        return d

    result: dict = {
        "stage1": _skipped(),
        "stage2": _skipped(),
        "stage3": _skipped(),
        "errors": [],
    }

    # ── Load openpyxl ─────────────────────────────────────────────────────────
    try:
        import openpyxl
    except ImportError:
        result["errors"].append("openpyxl not installed. Run: pip install openpyxl")
        return result

    path = Path(file_path)
    if not path.exists():
        result["errors"].append(f"File not found: {file_path}")
        return result

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        result["errors"].append(f"Cannot open Excel file: {e}")
        return result

    # ── Read sheets ───────────────────────────────────────────────────────────
    try:
        meas = _read_measurements_sheet(wb)
    except ValueError as e:
        result["errors"].append(str(e))
        return result

    try:
        temperatures, K_data = _read_thermal_sheet(wb)
    except ValueError as e:
        result["errors"].append(f"Thermal sheet error: {e}")
        temperatures, K_data = None, None

    # ── Build elastic targets ─────────────────────────────────────────────────
    targets_el: dict = {}
    sigmas_el: dict  = {}

    for excel_field, model_field in _ELASTIC_TO_MODEL.items():
        val = meas.get(excel_field)
        if val is not None and val != -1.0:
            targets_el[model_field] = float(val)

    for excel_field, model_field in _ELASTIC_SIGMA_TO_MODEL.items():
        val = meas.get(excel_field)
        if val is not None and val > 0.0 and model_field in targets_el:
            sigmas_el[model_field] = float(val)

    if not targets_el:
        result["errors"].append(
            "No elastic measurements in 'measurements' sheet. "
            "Provide at least E1_MPa and one of E2_MPa / G12_MPa."
        )
        return result

    # ── Load datasheet ────────────────────────────────────────────────────────
    try:
        datasheet = _smat.get_model_inputs(fiber_id, polymer_id)
    except Exception as e:
        result["errors"].append(f"Error loading material datasheets: {e}")
        return result

    # ── Build known microstructure ────────────────────────────────────────────
    known_micro: dict = {}
    if fiber_massfrac > 0:
        known_micro["fiber_massfrac"] = fiber_massfrac
    if aspect_ratio > 0:
        known_micro["ar"] = aspect_ratio
    if a11 > 0:
        known_micro["a11"] = a11
    if a22 > 0:
        known_micro["a22"] = a22
    known_micro.setdefault("a12", 0.0)
    known_micro.setdefault("a13", 0.0)
    known_micro.setdefault("a23", 0.0)

    for field in ("ar", "fiber_massfrac"):
        if field not in known_micro and field in datasheet:
            known_micro[field] = datasheet[field]

    has_shear = bool(set(targets_el.keys()) & _SHEAR_POISSON_MEASUREMENTS)
    base_free = _STAGE1_FREE if has_shear else _STAGE1_FREE_WITHOUT_POISSON
    free_vars = [v for v in base_free if v not in known_micro]

    free_set = set(free_vars)
    _model_fields_el = set(_sfwd.get_input_fields("elastic"))
    fixed_inputs_el = {k: v for k, v in datasheet.items() if k not in free_set and k in _model_fields_el}
    fixed_inputs_el.update({k: v for k, v in known_micro.items() if k in _model_fields_el})

    bounds    = {k: _DEFAULT_BOUNDS[k] for k in free_vars if k in _DEFAULT_BOUNDS}
    init_vals = []
    for k in free_vars:
        if k in _PROBLEM_JSON_INIT:
            init_vals.append(_PROBLEM_JSON_INIT[k])
        elif k in datasheet:
            init_vals.append(float(datasheet[k]))
        elif k in fixed_inputs_el:
            init_vals.append(float(fixed_inputs_el[k]))
        elif k in _DEFAULT_BOUNDS:
            lo, hi = _DEFAULT_BOUNDS[k]
            init_vals.append((lo + hi) / 2.0)
        else:
            init_vals.append(0.0)

    solver_cfg_el = dict(_ELASTIC_SOLVER_CFG)
    if sigmas_el:
        solver_cfg_el["use_epsilon_loss"] = True

    # ── Stage 1 ───────────────────────────────────────────────────────────────
    try:
        inv1 = _sinv.run_inverse(
            model_name="elastic",
            fixed_inputs=fixed_inputs_el,
            free_inputs=free_vars,
            bounds=bounds,
            target_outputs=targets_el,
            sigmas=sigmas_el or None,
            solver_cfg=solver_cfg_el,
            init_vals=init_vals,
        )
    except Exception as e:
        result["errors"].append(f"Stage 1 solver error: {e}")
        result["stage1"]["status"] = "fail"
        return result

    err1 = inv1["final_error"]
    orientation_warning1 = inv1.get("orientation_warning")
    stage1_pass = err1 < _STAGE1_POOR_FIT_THRESHOLD and not orientation_warning1
    result["stage1"] = {
        "status":    "pass" if stage1_pass else "fail",
        "fit_error": err1,
        "inferred":  dict(inv1["opt_free"]),
        "orientation_warning": orientation_warning1,
        "save_data": {
            "result": {
                "model":             "elastic",
                "opt_free":          inv1["opt_free"],
                "fixed_inputs":      fixed_inputs_el,
                "predicted_outputs": inv1["predicted_outputs"],
                "target_outputs":    inv1["target_outputs"],
                "sigmas":            sigmas_el,
                "final_error":       err1,
                "solver_cfg":        inv1["solver_cfg"],
            },
            "fiber_id":   fiber_id,
            "polymer_id": polymer_id,
            "printer_id": printer_id,
        },
    }

    if not stage1_pass:
        if orientation_warning1:
            result["errors"].append(f"Stage 1 orientation tensor invalid: {orientation_warning1}")
        if err1 >= _STAGE1_POOR_FIT_THRESHOLD:
            result["errors"].append(
                f"Stage 1 fit quality poor (fit_error={err1:.5f}). "
                "Check elastic measurements and material assignment."
            )
        result["stage2"] = _skipped("Skipped — Stage 1 failed.")
        result["stage3"] = _skipped("Skipped — Stage 1 failed.")
        return result

    # ── Stage 2: build inputs from Stage 1 in-memory (no DB round-trip) ──────
    targets_te: dict = {}
    sigmas_te: dict  = {}

    for excel_field, model_field in _CTE_TO_MODEL.items():
        val = meas.get(excel_field)
        if val is not None and val != -1.0:
            targets_te[model_field] = float(val)

    for excel_field, model_field in _CTE_SIGMA_TO_MODEL.items():
        val = meas.get(excel_field)
        if val is not None and val > 0.0 and model_field in targets_te:
            sigmas_te[model_field] = float(val)

    if "CTE11" not in targets_te or "CTE22" not in targets_te:
        result["stage2"] = _skipped("Skipped — CTE11_per_K and CTE22_per_K not in measurements sheet.")
        result["stage3"] = _skipped("Skipped — Stage 2 was skipped (no CTE measurements).")
        return result

    # Merge Stage 1 outputs with fixed inputs to get the full input set
    all_inputs_s1 = {**fixed_inputs_el, **inv1["opt_free"]}

    free_set_te = set(_TE_FREE)
    _te_fields  = set(_sfwd.get_input_fields("thermoelastic"))
    fixed_inputs_te = {
        k: v for k, v in all_inputs_s1.items()
        if k not in free_set_te and k in _te_fields
    }

    bounds_te    = {k: _DEFAULT_BOUNDS[k] for k in _TE_FREE if k in _DEFAULT_BOUNDS}
    init_vals_te = [_TE_INIT.get(k, 0.0) for k in _TE_FREE]

    solver_cfg_te = dict(_ELASTIC_SOLVER_CFG)
    if sigmas_te:
        solver_cfg_te["use_epsilon_loss"] = True

    try:
        inv2 = _sinv.run_inverse(
            model_name="thermoelastic",
            fixed_inputs=fixed_inputs_te,
            free_inputs=_TE_FREE,
            bounds=bounds_te,
            target_outputs=targets_te,
            sigmas=sigmas_te or None,
            solver_cfg=solver_cfg_te,
            init_vals=init_vals_te,
        )
    except Exception as e:
        result["errors"].append(f"Stage 2 solver error: {e}")
        result["stage2"]["status"] = "fail"
        return result

    err2 = inv2["final_error"]
    stage2_pass = err2 < _STAGE2_POOR_FIT_THRESHOLD
    result["stage2"] = {
        "status":    "pass" if stage2_pass else "fail",
        "fit_error": err2,
        "inferred":  dict(inv2["opt_free"]),
        "save_data": {
            "result": {
                "model":             "thermoelastic",
                "opt_free":          inv2["opt_free"],
                "fixed_inputs":      fixed_inputs_te,
                "predicted_outputs": inv2["predicted_outputs"],
                "target_outputs":    inv2["target_outputs"],
                "sigmas":            sigmas_te,
                "final_error":       err2,
                "solver_cfg":        inv2["solver_cfg"],
            },
            "fiber_id":   fiber_id,
            "polymer_id": polymer_id,
            "printer_id": printer_id,
        },
    }

    if not stage2_pass:
        result["errors"].append(
            f"Stage 2 fit quality poor (fit_error={err2:.5f}). Aborting Stage 3."
        )
        result["stage3"] = _skipped("Skipped — Stage 2 failed.")
        return result

    # ── Stage 3: build inputs from Stage 1 in-memory ──────────────────────────
    if temperatures is None:
        result["stage3"] = _skipped("Skipped — no 'thermal' sheet in Excel file.")
        return result

    # rho_f/rho_m are stripped from fixed_inputs_el (filtered to elastic model fields only),
    # so fall back to the raw datasheet which always has density under fiber_density/matrix_density.
    fixed_inputs_th = {
        "ar_f":  all_inputs_s1.get("ar",             all_inputs_s1.get("ar_f")),
        "w_f":   all_inputs_s1.get("fiber_massfrac",  all_inputs_s1.get("w_f")),
        "rho_f": (all_inputs_s1.get("rho_f")
                  or datasheet.get("rho_f")
                  or datasheet.get("fiber_density")),
        "rho_m": (all_inputs_s1.get("rho_m")
                  or datasheet.get("rho_m")
                  or datasheet.get("matrix_density")),
        "a11":   all_inputs_s1.get("a11"),
        "a22":   all_inputs_s1.get("a22"),
        "a12":   all_inputs_s1.get("a12", 0.0),
        "a13":   all_inputs_s1.get("a13", 0.0),
        "a23":   all_inputs_s1.get("a23", 0.0),
    }
    missing = [k for k, v in fixed_inputs_th.items() if v is None]
    if missing:
        result["errors"].append(f"Stage 3 missing fields from Stage 1: {missing}")
        result["stage3"]["status"] = "fail"
        return result

    try:
        fwd_model  = _sfwd.load_forward("thermal")
        predictor  = _ithermal.make_batched_predictor(fwd_model)
        best_params, best_loss = _ithermal.run_inverse_estimation(
            temperatures=temperatures,
            K_data=K_data,
            predictor=predictor,
            fixed_inputs=fixed_inputs_th,
            n_restarts=n_restarts,
        )
    except Exception as e:
        result["errors"].append(f"Stage 3 solver error: {e}")
        result["stage3"]["status"] = "fail"
        return result

    stage3_pass = best_loss < _STAGE3_POOR_FIT_THRESHOLD
    k_f1 = float(best_params.l2)
    k_f2 = float(best_params.l2 / best_params.t)
    result["stage3"] = {
        "status":    "pass" if stage3_pass else "fail",
        "fit_error": float(best_loss),
        "inferred": {
            "k_f1": k_f1,
            "k_f2": k_f2,
            "p1":   float(best_params.p1),
            "p2":   float(best_params.p2),
            "t":    float(best_params.t),
        },
        "save_data": {
            "result": {
                "best_params": best_params,
                "best_loss":   best_loss,
            },
            "fiber_id":   fiber_id,
            "polymer_id": polymer_id,
        },
    }

    return result
